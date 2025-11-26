import base64
import logging
from io import BytesIO
from openpyxl import load_workbook
from markupsafe import Markup

from odoo import api, fields, models, _
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

class EdiMessage(models.Model):
    _inherit = 'edi.message'

    def unpack(self):
        result = super().unpack()
        _logger.error("This works!!!"*10)
        if not result:
            _logger.error("Were in side two!!!"*10)
            file_extension = self.name.split(".")[-1]
            if file_extension == "xlsx":
               self.load_excel()
               result = True
            else:
               _logger.info(f"{self.name} is not a xlsx file")
        return result
    
    def load_excel(self):

        excel_file = BytesIO(base64.b64decode(self.payload))
        try:
            wb = load_workbook(excel_file)
        except Exception as e:
            raise UserError(f"The file given could not be read as an Excel file!\n\n{e}")
        
        header_fields = {
            "Bolag": "parent_id",
            "Namn": "name",
            "Land": "country_id",
            "Stad": "city",
            "Telefon": "phone",}

        vals_list = []

        for sheet in wb:
            first_row = next(sheet.iter_rows(values_only=True,min_row=1,max_row=1))
            word_index,filterd_row_fields = self.index_word_filter(first_row,header_fields)
            get_field_row_values = self.get_field_row_values(sheet,word_index,filterd_row_fields)
            vals_list.extend(get_field_row_values)
        
        self.create_records(vals_list)
        

    
    def create_records(self, vals_list):
        _logger.error(f"{vals_list=}")
        removed_prefix = self.name.split("-")[-1]
        split_filename = removed_prefix.split(".")
        model = self.env[split_filename[0].replace("_",".")]
        res_ids = self.env[model._name].create(vals_list)
        links = [
                    Markup('<a href="#" data-oe-model="res.partner" data-oe-id="{id}">{name}</a>').format(
                        id=res_id.id,
                        name=res_id.name if res_id.name else res_id.id
                    )
                    for res_id in res_ids
                ]
        move_links_str = Markup(", ").join(links)
        body = Markup("<p>Contact import completed successfully: {}</p>").format(move_links_str)
        self.message_post(body=body, message_type='comment')
        self.state = 'done'
        return res_ids

           
    # def create_account_move(self,header_fields,sheet,first_row):
    #     header_field_index,filterd_header_fields = self.index_word_filter(first_row,header_fields)
    #     header_field_values = self.get_header_field_values(sheet,header_field_index,filterd_header_fields)
    #     data = self._update_move_values(header_field_values)
    #     if journal_id := data.get("journal_id"):
    #           self._set_types(data,journal_id)
    #     _logger.error(f"{data=}")
    #     move_id = self.env["account.move"].create(data)
    #     return move_id
    
    # def create_account_move_line(self,header_field_rows,sheet,first_row,move_id):
    #     word_index,filterd_row_fields = self.index_word_filter(first_row,header_field_rows)
    #     header_field_row_values = self.get_header_field_row_values(sheet,word_index,filterd_row_fields)
    #     move_line_ids = []
    #     for data in header_field_row_values:
    #         data.update({"move_id":move_id.id})
    #         data = self._update_move_values(data)
    #         _logger.error(f"{data=}")
    #         move_line_ids.append(self.env["account.move.line"].create(data))
    #     return move_line_ids


    def is_relation_field(self, fname):
        field = self._fields.get(fname)
        return isinstance(field, (fields.Many2one, fields.One2many, fields.Many2many))

    def _update_values(self,data:dict):
        for key,value in data.items():
            _logger.error(f"{key=}")
            _logger.error(f"{value=}")
            new_value = value
            if key == "parent_id" and isinstance(value,str):
                new_value = self._is_string("res.partner",value)
            if key == "country_id" and isinstance(value,str):
                new_value = self._is_string("res.country",value)                           
            data.update({key:new_value})
        return data
    
    # def _set_types(self,data,journal_id):
    #     journal_id = self.env["account.journal"].browse(journal_id)
    #     if journal_id.type == "sale":
    #         data.update({
    #             "move_type": "out_invoice",
    #             }) 
    #     elif journal_id.type == "purchase":
    #         data.update({
    #             "move_type": "in_invoice",
    #             })
    #     return data
    
    # def is_account_code(self,value):
    #     new_value = value
    #     model_id = False
    #     if isinstance(value,str): 
    #         value = value.strip()
    #         ext_id = self.env.ref(value,raise_if_not_found=False)
    #         if ext_id:
    #             new_value = ext_id.id
    #     else:
    #         model_id = self.env["account.account"].search([
    #             "|",
    #             ("name","ilike",value),
    #             ("code","=",value)],
    #             limit=1)
    #     if model_id:
    #         new_value = model_id.id
    #     return new_value

    def _is_string(self,model,value):
        value = value.strip()
        new_value = value
        ext_id = self.env.ref(value,raise_if_not_found=False)
        model_id = self.env[model].search([
            "|",
            ("name","ilike",value),
            ("display_name","=",value)],
            limit=1)
        if ext_id:
            new_value = ext_id.id
        elif model_id:
            new_value = model_id.id
        return new_value

    def get_field_row_values(self,sheet,word_index,filterd_row_fields):
        datas = []
        for row in sheet.iter_rows(values_only=True,min_row=2,max_row=100,max_col=100):
            if not all(r is None for r in row):
                row_values = [row[i] for i in word_index]
                new_row = dict(zip(filterd_row_fields,row_values))
                new_row = self._update_values(new_row)
                datas.append(new_row)
        return datas

    def index_word_filter(self,first_row,words):
        word_index = []
        filterd_words_fields = []
        for word in words.keys():
            if word in first_row:
                row_index = first_row.index(word)
                word_index.append(row_index)
                filterd_words_fields.append(words[word])
        return word_index, filterd_words_fields
    
    